import json
import os
import sys
import tempfile
import types
import unittest
from pathlib import Path


def _install_test_stubs():
    logger = types.SimpleNamespace(
        info=lambda *args, **kwargs: None,
        warning=lambda *args, **kwargs: None,
        error=lambda *args, **kwargs: None,
        debug=lambda *args, **kwargs: None,
        exception=lambda *args, **kwargs: None,
    )

    config_module = types.ModuleType("config")
    log_config_module = types.ModuleType("config.log_config")
    log_config_module.app_logger = logger
    config_module.log_config = log_config_module
    sys.modules.setdefault("config", config_module)
    sys.modules["config.log_config"] = log_config_module


_install_test_stubs()

from openpyxl import Workbook, load_workbook

from core.pipelines.excel_translation_pipeline import write_translated_content_to_excel
from core.pipelines.subtitle_translation_pipeline import write_translated_content_to_srt
from core.pipelines.txt_translation_pipeline import write_translated_content_to_txt
from core.engine.writeback_same_ok import build_writeback_translation_map


class SameOkWritebackTests(unittest.TestCase):
    def test_build_writeback_translation_map_excludes_same_ok_items(self):
        translated_data = [
            {"count_src": 1, "translated": "DO NOT WRITE", "same_ok": True},
            {"count_src": 2, "translated": "你好", "same_ok": False},
        ]

        translations = build_writeback_translation_map(translated_data)

        self.assertNotIn("1", translations)
        self.assertEqual(translations["2"], "你好")

    @unittest.skip("待接入:上游 pipeline 尚未调用 build_writeback_translation_map")
    def test_txt_writeback_keeps_original_for_same_ok_items(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_file = root / "input.txt"
            source_file.write_text("", encoding="utf-8")

            temp_input_dir = root / "temp" / "input"
            temp_input_dir.mkdir(parents=True)
            (temp_input_dir / "all_content.json").write_text(
                json.dumps(
                    [
                        {"count_src": 1, "value": "A-101", "needs_translation": True},
                        {"count_src": 2, "value": "Hello", "needs_translation": True},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            translated_path = root / "translated.json"
            translated_path.write_text(
                json.dumps(
                    [
                        {"count_src": 1, "translated": "DO NOT WRITE", "same_ok": True},
                        {"count_src": 2, "translated": "你好"},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result_path = write_translated_content_to_txt(
                str(source_file),
                str(root / "src.json"),
                str(translated_path),
                str(root / "temp"),
                str(root / "result"),
                src_lang="en",
                dst_lang="zh",
            )

            self.assertEqual(Path(result_path).read_text(encoding="utf-8"), "A-101\n\n你好\n\n")

    @unittest.skip("待接入:上游 pipeline 尚未调用 build_writeback_translation_map")
    def test_srt_writeback_keeps_original_for_same_ok_items(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_file = root / "movie.srt"
            source_file.write_text("", encoding="utf-8")

            original_path = root / "src.json"
            original_path.write_text(
                json.dumps(
                    [
                        {
                            "count_src": 1,
                            "start_time": "00:00:01,000",
                            "end_time": "00:00:02,000",
                            "value": "A-101",
                        },
                        {
                            "count_src": 2,
                            "start_time": "00:00:03,000",
                            "end_time": "00:00:04,000",
                            "value": "Hello",
                        },
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            translated_path = root / "translated.json"
            translated_path.write_text(
                json.dumps(
                    [
                        {"count_src": 1, "translated": "DO NOT WRITE", "same_ok": True},
                        {"count_src": 2, "translated": "你好"},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result_path = write_translated_content_to_srt(
                str(source_file),
                str(original_path),
                str(translated_path),
                str(root / "result"),
                src_lang="en",
                dst_lang="zh",
            )

            self.assertIn(
                "1\n00:00:01,000 --> 00:00:02,000\nA-101\n\n",
                Path(result_path).read_text(encoding="utf-8"),
            )

    @unittest.skip("待接入:上游 pipeline 尚未调用 build_writeback_translation_map")
    def test_excel_writeback_keeps_original_cell_for_same_ok_items(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "book.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet["A1"] = "A-101"
            sheet["A2"] = "Hello"
            workbook.save(workbook_path)

            original_path = root / "src.json"
            original_path.write_text(
                json.dumps(
                    [
                        {"count_src": 1, "type": "cell", "sheet": "Sheet1", "row": 1, "column": 1, "value": "A-101"},
                        {"count_src": 2, "type": "cell", "sheet": "Sheet1", "row": 2, "column": 1, "value": "Hello"},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            translated_path = root / "translated.json"
            translated_path.write_text(
                json.dumps(
                    [
                        {"count_src": 1, "translated": "DO NOT WRITE", "same_ok": True},
                        {"count_src": 2, "translated": "你好"},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result_path = write_translated_content_to_excel(
                str(workbook_path),
                str(original_path),
                str(translated_path),
                str(root / "result"),
                src_lang="en",
                dst_lang="zh",
            )

            result_workbook = load_workbook(result_path)
            self.assertEqual(result_workbook["Sheet1"]["A1"].value, "A-101")
            self.assertEqual(result_workbook["Sheet1"]["A2"].value, "你好")


if __name__ == "__main__":
    unittest.main()
