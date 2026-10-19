# A model reply can carry characters that are not legal in XML 1.0 — a raw C0
# control (one real reply had \x03 where the source had Wingdings glyphs) or a
# lone surrogate. They are harmless in JSON, but write-back assigns them to lxml
# element text and serializes, which raises "All strings must be XML compatible"
# and loses a fully paid translation. They are stripped at the boundary.
import json
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from core.engine.xml_sanitize import (
    INVALID_XML_CHARS, strip_xml_invalid, strip_xml_invalid_in_json)
from core.pipelines.word_translation_pipeline import (
    extract_word_content_to_json, write_translated_content_to_word)


def test_strips_control_characters_but_keeps_xml_whitespace():
    assert strip_xml_invalid("a\x03b\x0bc\x0cd\x1fe") == "abcde"
    assert strip_xml_invalid("a\tb\nc\rd") == "a\tb\nc\rd"
    assert strip_xml_invalid("a\x00b") == "ab"


def test_keeps_private_use_codepoints():
    # Wingdings glyphs (U+F0xx) and the placeholder-mask sentinels (U+E000/1)
    # are legal XML and load-bearing: stripping them would break bullets and
    # mask round-tripping.
    glyphs = "".join(map(chr, (0xF063, 0xF03A, 0xF03E)))
    assert strip_xml_invalid(glyphs) == glyphs
    assert strip_xml_invalid("".join(map(chr, (0xE000, 0xE001)))) == \
        "".join(map(chr, (0xE000, 0xE001)))


def test_strips_lone_surrogates_and_non_characters():
    assert strip_xml_invalid("a\ud800b") == "ab"
    assert strip_xml_invalid("a" + chr(0xFFFE) + chr(0xFFFF) + "b") == "ab"


def test_non_strings_pass_through():
    for value in (None, 7, ["a"], {"a": 1}):
        assert strip_xml_invalid(value) == value


def test_is_idempotent():
    once = strip_xml_invalid("a\x03b\x00c")
    assert strip_xml_invalid(once) == once == "abc"


def _poisoned_json(tmp_path):
    path = str(tmp_path / "dst.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump([
            {"count_src": 1, "translated": "\x03\xba> 延迟时间已设置。"},
            {"count_src": 2, "translated": "clean text"},
        ], f, ensure_ascii=False, indent=4)
    return path


def test_json_sanitizer_cleans_only_the_translated_field(tmp_path):
    path = _poisoned_json(tmp_path)
    assert strip_xml_invalid_in_json(path) == 1
    with open(path, encoding="utf-8") as f:
        items = json.load(f)
    assert items[0]["translated"] == "\xba> 延迟时间已设置。"
    assert items[1]["translated"] == "clean text"


def test_json_sanitizer_leaves_a_clean_file_byte_identical(tmp_path):
    path = _poisoned_json(tmp_path)
    strip_xml_invalid_in_json(path)
    before = open(path, "rb").read()
    assert strip_xml_invalid_in_json(path) == 0
    assert open(path, "rb").read() == before


def _write_with(tmp_path, translated):
    src = str(tmp_path / "in.docx")
    doc = Document()
    doc.add_paragraph("Die Verzögerungszeit ist eingestellt.")
    doc.save(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": i["count_src"], "translated": translated}
                   for i in items], f, ensure_ascii=False)
    return src, src_json, dst_json


def test_control_character_no_longer_aborts_writeback(tmp_path):
    # End-to-end for the original crash: the sanitizer runs between restore and
    # write, so a poisoned result file still produces a valid document.
    src, src_json, dst_json = _write_with(tmp_path, "\x03\xba> 延迟时间已设置。")
    assert strip_xml_invalid_in_json(dst_json) == 1

    result = write_translated_content_to_word(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        src_lang="de", dst_lang="zh")

    xml = zipfile.ZipFile(result).read("word/document.xml").decode("utf-8")
    assert not INVALID_XML_CHARS.search(xml)
    assert Document(result).paragraphs[0].text == "\xba> 延迟时间已设置。"
