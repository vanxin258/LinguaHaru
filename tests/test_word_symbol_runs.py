# Word "symbol font" runs hold font glyph indices in w:t rather than text: a
# Wingdings bullet is literally U+F063 U+F03A U+F03E, because those private-use
# codepoints map to glyphs in the font. Such a run must never be translated —
# the model either drops the glyphs (the bullet disappears from the output) or
# mangles them (one real reply carried a raw \x03, which lxml refuses to
# serialize and which aborted a whole document's write-back) — and translated
# text that lands in the run inherits the symbol font and renders as garbage.
# They are kept verbatim behind a {{SYM_n}} placeholder, like super/subscript
# markers. Deterministic: a fake translation map stands in for the LLM.
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from core.pipelines.word_translation_pipeline import (
    extract_word_content_to_json, write_translated_content_to_word)

# U+F063 U+F03A U+F03E: the Wingdings glyph indices for the bullet, stored
# literally in the run's w:t. Built from code points so the test file keeps no
# invisible private-use characters.
GLYPH = "".join(map(chr, (0xF063, 0xF03A, 0xF03E)))
SOURCE = "Die Verzögerungszeit ist eingestellt."
CHINESE = "延迟时间已设置。"


def _build(path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run(GLYPH).font.name = "Wingdings"
    p.add_run(" " + SOURCE)
    doc.save(path)


def _extract(tmp_path):
    src = str(tmp_path / "in.docx")
    _build(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    return src, src_json, items


def _write(tmp_path, src, src_json, items, translated):
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": i["count_src"], "translated": translated}
                   for i in items], f, ensure_ascii=False)
    result = write_translated_content_to_word(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        src_lang="de", dst_lang="zh")
    return Document(result)


def test_symbol_glyph_run_is_masked(tmp_path):
    _, _, items = _extract(tmp_path)
    assert items[0]["value"] == "{{SYM_1}} " + SOURCE


def test_dropped_placeholder_does_not_leak_the_symbol_font(tmp_path):
    # If the reply lost the placeholder there is nothing to anchor the glyph run
    # to, so the bullet is not restored here. What must NOT happen is the
    # translation being written into the Wingdings run - that is what made the
    # Chinese render in the symbol font. Losing the bullet outright is caught
    # upstream instead (see the structural-break test below).
    src, src_json, items = _extract(tmp_path)
    para = _write(tmp_path, src, src_json, items, CHINESE).paragraphs[0]

    assert "".join(r.text for r in para.runs) == CHINESE
    assert all(r.font.name != "Wingdings" for r in para.runs)


def test_dropped_symbol_marker_is_a_structural_break():
    # The marker is validated like every other {{...}} placeholder, so a reply
    # that drops it falls back to the source text rather than shipping a
    # document with the bullet silently missing.
    from core.engine.translation_checker import _structural_intact

    assert _structural_intact("{{SYM_1}} " + SOURCE, "{{SYM_1}} " + CHINESE)
    assert not _structural_intact("{{SYM_1}} " + SOURCE, CHINESE)


def test_glyph_restored_when_model_keeps_the_placeholder(tmp_path):
    src, src_json, items = _extract(tmp_path)
    para = _write(tmp_path, src, src_json, items,
                  "{{SYM_1}} " + CHINESE).paragraphs[0]

    assert para.runs[0].text == GLYPH
    assert para.runs[0].font.name == "Wingdings"
    assert "".join(r.text for r in para.runs[1:]) == " " + CHINESE
    assert "{{" not in para._p.xml


def test_symbol_run_never_donates_its_formatting(tmp_path):
    # A run set in a symbol font must not be picked as the paragraph's
    # formatting template either (the `(SET)` case: latin text came back inside
    # a Wingdings run). The glyph run is the FIRST run, so it is the one the old
    # fallback would have chosen.
    src, src_json, items = _extract(tmp_path)
    para = _write(tmp_path, src, src_json, items, CHINESE).paragraphs[0]
    for run in para.runs[1:]:
        assert run.font.name != "Wingdings"
        assert run._r.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr"
        ) is not None


def test_plain_text_paragraph_is_untouched_by_symbol_handling(tmp_path):
    # Guard against over-matching: an ordinary paragraph must not gain markers.
    doc = Document()
    doc.add_paragraph("Anzeige/Einstellen Soll-Werte (SET) für Temperatur")
    src = str(tmp_path / "plain.docx")
    doc.save(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    assert items[0]["value"] == "Anzeige/Einstellen Soll-Werte (SET) für Temperatur"
